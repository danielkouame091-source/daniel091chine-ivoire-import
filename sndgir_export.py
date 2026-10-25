"""Exports Excel et PDF de n'importe quel tableau de l'application.

Sécurité : une cellule de texte qui commencerait par = + - @ serait exécutée comme une formule par Excel
(injection de formule) ; on la neutralise. Limite PDF : 2 000 lignes, texte tronqué à 40 caractères.
"""
from __future__ import annotations

import io
import re
from datetime import datetime

import pandas as pd
import streamlit as st

from sndgir_i18n import T, trad_df
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

PDF_MAX_LIGNES = 2000
PDF_MAX_CAR = 40
_DEBUT_FORMULE = re.compile(r"^[=+\-@\t\r]")


def _neutraliser(valeur):
    if isinstance(valeur, str) and _DEBUT_FORMULE.match(valeur):
        return "'" + valeur
    return valeur


def _nom_feuille(titre: str) -> str:
    return re.sub(r"[\[\]\*\?/\\:]", " ", titre).strip()[:31] or "Export"


def df_vers_excel(df: pd.DataFrame, titre: str = "Export") -> bytes:
    df, titre = trad_df(df), T(titre)
    propre = df.copy()
    for col in propre.columns:
        if not pd.api.types.is_numeric_dtype(propre[col]) and not pd.api.types.is_datetime64_any_dtype(propre[col]):
            propre[col] = propre[col].map(_neutraliser)
    tampon = io.BytesIO()
    with pd.ExcelWriter(tampon, engine="openpyxl") as w:
        propre.to_excel(w, index=False, sheet_name=_nom_feuille(titre))
        feuille = w.sheets[_nom_feuille(titre)]
        for cellule in feuille[1]:
            cellule.font = Font(bold=True, color="FFFFFF")
            cellule.fill = PatternFill("solid", fgColor="1F3A5F")
            cellule.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, col in enumerate(propre.columns, start=1):
            largeur = max([len(str(col))] + [len(str(v)) for v in propre[col].head(200)]) + 2
            feuille.column_dimensions[get_column_letter(i)].width = min(max(largeur, 10), 45)
        feuille.freeze_panes = "A2"
    return tampon.getvalue()


def _texte_pdf(valeur) -> str:
    if valeur is None or (isinstance(valeur, float) and pd.isna(valeur)):
        return ""
    if isinstance(valeur, float):
        texte = f"{valeur:,.2f}".replace(",", " ").rstrip("0").rstrip(".") if abs(valeur) < 1e15 else str(valeur)
    else:
        texte = str(valeur)
    texte = texte.encode("latin-1", "replace").decode("latin-1")     # pas d'emoji : police Helvetica
    return texte if len(texte) <= PDF_MAX_CAR else texte[:PDF_MAX_CAR - 1] + "…".encode("latin-1", "replace").decode("latin-1")


def df_vers_pdf(df: pd.DataFrame, titre: str = "Export", entreprise: str = "") -> bytes:
    df, titre = trad_df(df), T(titre)
    tampon = io.BytesIO()
    doc = SimpleDocTemplate(tampon, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=titre)
    styles = getSampleStyleSheet()
    entete = [Paragraph(_texte_pdf(titre).replace("&", "&amp;").replace("<", "&lt;"), styles["Title"])]
    sous = f"{entreprise} — édité le {datetime.now():%d/%m/%Y %H:%M} — {len(df)} ligne(s)"
    entete.append(Paragraph(_texte_pdf(sous).replace("&", "&amp;").replace("<", "&lt;"), styles["Normal"]))
    if len(df) > PDF_MAX_LIGNES:
        entete.append(Paragraph(f"Seules les {PDF_MAX_LIGNES} premières lignes sont imprimées : utilisez l'export Excel.", styles["Normal"]))
    entete.append(Spacer(1, 6 * mm))
    corps = [[_texte_pdf(c) for c in df.columns]] + [[_texte_pdf(v) for v in ligne] for ligne in df.head(PDF_MAX_LIGNES).itertuples(index=False)]
    if len(df.columns) == 0:
        corps = [["(aucune colonne)"]]
    t = Table(corps, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F4F9")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C5CCD8")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
    doc.build(entete + [t])
    return tampon.getvalue()


def bouton_export(df: pd.DataFrame, nom: str, titre: str | None = None) -> None:
    """Deux boutons de téléchargement (Excel, PDF) sous un tableau. Rien si le tableau est vide."""
    if df is None or df.empty:
        return
    titre = titre or nom
    entreprise = st.session_state.get("entreprise_nom", "")
    cle = re.sub(r"\W+", "_", nom)
    horodatage = datetime.now().strftime("%Y%m%d_%H%M")
    c1, c2, _ = st.columns([1, 1, 4])
    c1.download_button("⬇️ Excel", df_vers_excel(df, titre), f"{cle}_{horodatage}.xlsx",
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key=f"xl_{cle}")
    c2.download_button("⬇️ PDF", df_vers_pdf(df, titre, entreprise), f"{cle}_{horodatage}.pdf",
                       "application/pdf", key=f"pdf_{cle}")
