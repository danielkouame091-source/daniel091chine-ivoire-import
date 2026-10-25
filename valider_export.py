"""Valide les exports Excel et PDF (python3 valider_export.py)."""
import io, sys
from unittest.mock import MagicMock
import pandas as pd
sys.modules.setdefault("streamlit", MagicMock())
sys.path.insert(0, ".")
import sndgir_export as e
from openpyxl import load_workbook
df = pd.DataFrame({"Client": ["SARL Ivoir'Tech", "=CMD|' /C calc'!A0", "+1+1", "Groupe Bâtir SA"], "Montant": [14000000.0, 2.5, -3.0, None],
                   "Note": ["émoji 😀 é", "x" * 200, "ok", "@SUM(A1)"]})
x = e.df_vers_excel(df, "Dossiers / janvier [test]?")
ws = load_workbook(io.BytesIO(x)).active
assert ws.title == "Dossiers   janvier  test" or len(ws.title) <= 31, ws.title
assert [c.value for c in ws[1]] == ["Client", "Montant", "Note"]
vals = [ws.cell(r, 1).value for r in range(2, 6)]
assert vals[0] == "SARL Ivoir'Tech" and vals[1].startswith("'=") and vals[2].startswith("'+") and ws.cell(5, 3).value.startswith("'@")
assert ws.cell(2, 2).value == 14000000.0 and ws.freeze_panes == "A2"
assert all(not str(c.value).startswith("=") for row in ws.iter_rows(min_row=2) for c in row if c.value is not None)
p = e.df_vers_pdf(df, "Dossiers <b>& co</b>", "Cabinet Excellence")
assert p[:5] == b"%PDF-" and len(p) > 1500
grand = pd.DataFrame({"a": range(5000), "b": ["texte"] * 5000})
pg = e.df_vers_pdf(grand, "Gros tableau")
assert pg[:5] == b"%PDF-" and pg.count(b"/Type /Page\n") + pg.count(b"/Type /Page ") >= 2
assert e.df_vers_pdf(pd.DataFrame(), "Vide")[:5] == b"%PDF-" and e.df_vers_excel(pd.DataFrame({"a": []}), "Vide")[:2] == b"PK"
assert e._texte_pdf(None) == "" and e._texte_pdf(float("nan")) == "" and e._texte_pdf(1234567.5) == "1 234 567.5"
e.bouton_export(pd.DataFrame(), "x")                       # ne plante pas sur un tableau vide
print("EXPORT OK")
